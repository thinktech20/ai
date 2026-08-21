from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

import pandas as pd
from openpyxl import load_workbook

from .config import clean_scalar, normalize_issue_name


SCORING_WORKBOOK_NAME = "Scoring Summary and Follow-up Items.xlsx"
GROUND_TRUTH_CSV_NAME = "ground_truth.csv"
DIFF_CSV_NAME = "ground_truth_diff.csv"

RISK_TO_SCORE = {
    "Not Mentioned": 1,
    "Light": 1,
    "Med": 2,
    "Heavy": 3,
    "IMMEDIATE ACTION": 4,
}

# Map ground-truth labels from the wide-format CSV to canonical risk labels
_GT_RISK_ALIASES = {
    "light": "Light",
    "medium": "Med",
    "med": "Med",
    "heavy": "Heavy",
    "immediate action": "IMMEDIATE ACTION",
    "no data": "Not Mentioned",
    "no data or medium": "Med",       # conservative: treat ambiguous as higher
    "l or n/a": "Light",
    "light or medium": "Med",          # conservative: treat ambiguous as higher
    "medium or high": "Heavy",         # conservative: treat ambiguous as higher
}


def _normalize_risk_label(value: Any) -> str:
    text = clean_scalar(value).rstrip("* ").strip()
    if not text:
        return ""
    # Try alias lookup (case-insensitive)
    canonical = _GT_RISK_ALIASES.get(text.lower())
    if canonical:
        return canonical
    return text


def _select_scoring_sheet(workbook) -> Any:
    for sheet_name in workbook.sheetnames:
        if normalize_issue_name(sheet_name) != normalize_issue_name("Follow-Ups"):
            return workbook[sheet_name]
    return workbook[workbook.sheetnames[0]]


def _iter_serial_columns(worksheet) -> list[tuple[str, int, Optional[int]]]:
    serial_columns: list[tuple[str, int, Optional[int]]] = []
    for column_index in range(2, worksheet.max_column + 1, 2):
        serial_number = clean_scalar(worksheet.cell(1, column_index).value)
        if not serial_number:
            continue
        rationale_column = column_index + 1 if column_index + 1 <= worksheet.max_column else None
        serial_columns.append((serial_number, column_index, rationale_column))
    return serial_columns


def parse_scoring_workbook(workbook_path: str | Path) -> pd.DataFrame:
    workbook_file = Path(workbook_path)
    if not workbook_file.exists():
        raise FileNotFoundError(f"Scoring workbook not found: {workbook_file}")

    workbook = load_workbook(workbook_file, data_only=True)
    worksheet = _select_scoring_sheet(workbook)
    serial_columns = _iter_serial_columns(worksheet)

    records: list[dict[str, Any]] = []
    for row_index in range(2, worksheet.max_row + 1):
        issue_name = clean_scalar(worksheet.cell(row_index, 1).value)
        if not issue_name:
            continue

        normalized_issue = normalize_issue_name(issue_name)
        for serial_number, severity_column, rationale_column in serial_columns:
            expected_risk = _normalize_risk_label(
                worksheet.cell(row_index, severity_column).value
            )
            if not expected_risk:
                continue

            rationale = ""
            if rationale_column is not None:
                rationale = clean_scalar(worksheet.cell(row_index, rationale_column).value)

            records.append(
                {
                    "sheet_name": worksheet.title,
                    "source_row": row_index,
                    "serial_number": serial_number,
                    "issue_name": issue_name,
                    "normalized_issue_name": normalized_issue,
                    "expected_risk": expected_risk,
                    "rationale": rationale,
                }
            )

    if not records:
        raise ValueError(f"No scoring records found in workbook: {workbook_file}")

    df = pd.DataFrame(records)
    df["occurrence_index"] = df.groupby("normalized_issue_name").cumcount() + 1
    return df


def _detect_rotor_start_row(scoring_df: pd.DataFrame) -> Optional[int]:
    ir_pi_key = normalize_issue_name("IR, PI")
    matching_rows = scoring_df.loc[
        scoring_df["normalized_issue_name"].eq(ir_pi_key), "source_row"
    ].drop_duplicates().sort_values()
    if len(matching_rows) >= 2:
        return int(matching_rows.iloc[1])
    return None


def _detect_contiguous_end_row(scoring_df: pd.DataFrame, start_row: int) -> int:
    candidate_rows = (
        scoring_df.loc[scoring_df["source_row"] >= start_row, "source_row"]
        .drop_duplicates()
        .sort_values()
        .tolist()
    )
    contiguous_end = start_row
    previous_row = start_row
    for row in candidate_rows[1:]:
        if int(row) != previous_row + 1:
            break
        contiguous_end = int(row)
        previous_row = int(row)
    return contiguous_end


def build_ground_truth_df(workbook_path: str | Path) -> pd.DataFrame:
    scoring_df = parse_scoring_workbook(workbook_path)
    rotor_start_row = _detect_rotor_start_row(scoring_df)
    working = scoring_df.copy()
    if rotor_start_row is not None:
        rotor_end_row = _detect_contiguous_end_row(scoring_df, rotor_start_row)
        working = working.loc[
            working["source_row"].between(rotor_start_row, rotor_end_row)
        ].copy()

    working = working.sort_values(["source_row", "serial_number"]).drop_duplicates(
        ["serial_number", "normalized_issue_name"], keep="last"
    )
    working["expected_score"] = working["expected_risk"].map(RISK_TO_SCORE)

    unknown_risks = sorted(
        working.loc[working["expected_score"].isna(), "expected_risk"].dropna().unique().tolist()
    )
    if unknown_risks:
        raise ValueError(f"Unknown expected risk values encountered: {unknown_risks}")

    ordered_columns = [
        "serial_number",
        "issue_name",
        "normalized_issue_name",
        "expected_risk",
        "expected_score",
        "rationale",
        "sheet_name",
        "source_row",
        "occurrence_index",
    ]
    return working.loc[:, ordered_columns].reset_index(drop=True)


def _load_wide_ground_truth(csv_path: str | Path) -> pd.DataFrame:
    """Load the wide-format ground truth CSV (serials as columns) and melt to long format.

    Expected columns: component, issue_name, <blank>, serial1, serial2, ...
    Returns DataFrame with: serial_number, issue_name, normalized_issue_name,
                            expected_risk, expected_score, component
    """
    df = pd.read_csv(csv_path).fillna("")
    # Drop the blank column (Unnamed: 2)
    df = df.loc[:, ~df.columns.str.startswith("Unnamed")]

    # Forward-fill component for rows where it's blank
    if "component" in df.columns:
        df["component"] = df["component"].replace("", pd.NA).ffill().fillna("")

    serial_cols = [c for c in df.columns if c not in ("component", "issue_name")]

    records: list[dict[str, Any]] = []
    for _, row in df.iterrows():
        issue_name = clean_scalar(row["issue_name"])
        if not issue_name:
            continue
        component = clean_scalar(row.get("component", ""))
        for serial in serial_cols:
            raw_risk = clean_scalar(row[serial])
            expected_risk = _normalize_risk_label(raw_risk)
            if not expected_risk:
                continue
            records.append({
                "serial_number": serial.strip(),
                "issue_name": issue_name,
                "normalized_issue_name": normalize_issue_name(issue_name),
                "expected_risk": expected_risk,
                "expected_score": RISK_TO_SCORE.get(expected_risk),
                "component": component,
                "rationale": "",
                "sheet_name": "",
                "source_row": 0,
                "occurrence_index": 0,
            })

    if not records:
        raise ValueError(f"No ground truth records found in: {csv_path}")

    result = pd.DataFrame(records)
    unknown = sorted(
        result.loc[result["expected_score"].isna(), "expected_risk"].dropna().unique().tolist()
    )
    if unknown:
        print(f"[Eval] WARNING: Unknown risk labels in ground truth (unmapped): {unknown}")

    return result


def _is_wide_format(csv_path: str | Path) -> bool:
    """Detect if a ground truth CSV is wide format (serials as columns)."""
    df = pd.read_csv(csv_path, nrows=0)
    cols = [c.lower().strip() for c in df.columns]
    # Wide format has 'component' and 'issue_name' as first two columns
    # Long format has 'serial_number' as first column
    return "component" in cols and "serial_number" not in cols


def ensure_ground_truth_csv(project_dir: str | Path) -> Optional[Path]:
    project_root = Path(project_dir)
    workbook_path = project_root / SCORING_WORKBOOK_NAME
    ground_truth_csv = project_root / GROUND_TRUTH_CSV_NAME

    if workbook_path.exists():
        ground_truth_df = build_ground_truth_df(workbook_path)
        ground_truth_df.to_csv(ground_truth_csv, index=False, encoding="utf-8")
        return ground_truth_csv

    if ground_truth_csv.exists():
        return ground_truth_csv

    return None


def load_actual_condensed(condensed_csv_path: str | Path) -> pd.DataFrame:
    condensed_csv = Path(condensed_csv_path)
    if not condensed_csv.exists():
        raise FileNotFoundError(f"Condensed CSV not found: {condensed_csv}")

    wide_df = pd.read_csv(condensed_csv).fillna("")
    issue_column = wide_df.columns[0]
    wide_df = wide_df.rename(columns={issue_column: "issue_name"})

    records: list[dict[str, Any]] = []
    for _, row in wide_df.iterrows():
        issue_name = clean_scalar(row["issue_name"])
        normalized_issue = normalize_issue_name(issue_name)
        for serial_number in wide_df.columns[1:]:
            actual_risk = _normalize_risk_label(clean_scalar(row[serial_number]))
            records.append(
                {
                    "serial_number": clean_scalar(serial_number),
                    "actual_issue_name": issue_name,
                    "normalized_issue_name": normalized_issue,
                    "actual_risk": actual_risk,
                }
            )

    actual_df = pd.DataFrame(records)
    actual_df["actual_score"] = actual_df["actual_risk"].map(RISK_TO_SCORE)
    unknown_risks = sorted(
        actual_df.loc[
            actual_df["actual_risk"].ne("") & actual_df["actual_score"].isna(),
            "actual_risk",
        ].unique().tolist()
    )
    if unknown_risks:
        raise ValueError(f"Unknown actual risk values encountered: {unknown_risks}")
    return actual_df


def build_diff_df(
    condensed_csv_path: str | Path,
    ground_truth_csv_path: str | Path,
) -> pd.DataFrame:
    gt_path = Path(ground_truth_csv_path)
    if _is_wide_format(gt_path):
        ground_truth_df = _load_wide_ground_truth(gt_path)
    else:
        ground_truth_df = pd.read_csv(gt_path).fillna("")
    actual_df = load_actual_condensed(condensed_csv_path)

    diff_df = ground_truth_df.merge(
        actual_df,
        on=["serial_number", "normalized_issue_name"],
        how="left",
    )
    diff_df["score_diff_actual_minus_expected"] = (
        diff_df["actual_score"] - diff_df["expected_score"]
    )

    def classify(score_diff: Any) -> str:
        if pd.isna(score_diff):
            return "missing_actual"
        if float(score_diff) == 0:
            return "correct"
        if float(score_diff) > 0:
            return "false_positive"
        return "false_negative"

    diff_df["classification"] = diff_df["score_diff_actual_minus_expected"].map(classify)
    diff_df["label_match"] = diff_df["actual_risk"].eq(diff_df["expected_risk"])

    ordered_columns = [
        "serial_number",
        "issue_name",
        "normalized_issue_name",
        "expected_risk",
        "expected_score",
        "actual_issue_name",
        "actual_risk",
        "actual_score",
        "score_diff_actual_minus_expected",
        "classification",
        "label_match",
        "rationale",
        "sheet_name",
        "source_row",
        "occurrence_index",
    ]
    return diff_df.loc[:, ordered_columns].reset_index(drop=True)


def summarize_diff(diff_df: pd.DataFrame) -> dict[str, Any]:
    valid = diff_df.loc[diff_df["classification"].ne("missing_actual")].copy()
    score_diff = valid["score_diff_actual_minus_expected"]
    return {
        "rows_compared": int(len(valid)),
        "correct": int(score_diff.eq(0).sum()),
        "false_positive": int(score_diff.gt(0).sum()),
        "false_negative": int(score_diff.lt(0).sum()),
        "missing_actual": int(diff_df["classification"].eq("missing_actual").sum()),
        "score_mapping": dict(RISK_TO_SCORE),
        "diff_definition": "actual_score - expected_score",
    }


def evaluate_condensed_against_ground_truth(
    condensed_csv_path: str | Path,
    ground_truth_csv_path: str | Path,
    diff_csv_path: str | Path,
) -> dict[str, Any]:
    diff_df = build_diff_df(condensed_csv_path, ground_truth_csv_path)
    diff_csv = Path(diff_csv_path)
    diff_csv.parent.mkdir(parents=True, exist_ok=True)
    diff_df.to_csv(diff_csv, index=False, encoding="utf-8")

    return {
        "ground_truth_csv": str(Path(ground_truth_csv_path)),
        "diff_csv": str(diff_csv),
        "stats": summarize_diff(diff_df),
    }